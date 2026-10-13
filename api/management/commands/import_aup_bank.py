"""Import the AUP question bank (the manager's Excel) into
api/data/aup_question_bank.json, which aup_builder.py reads at runtime.

    python manage.py import_aup_bank "E:\\AUP_Questionnaire_Question_Bank_with_wording.xlsx"

Run it again whenever the workbook changes, then commit the JSON. Needs
openpyxl only for this command (pip install openpyxl) — not at runtime.
"""

import json
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

OUT = Path(__file__).resolve().parents[2] / "data" / "aup_question_bank.json"
TYPES = {"Single select": "single", "Multi select": "multi", "Yes/No": "single",
         "Text": "text", "Email": "email", "Date": "date"}


def _s(value) -> str:
    return "" if value is None else str(value).strip()


class Command(BaseCommand):
    help = "Import the AUP questionnaire workbook into api/data/aup_question_bank.json."

    def add_arguments(self, parser):
        parser.add_argument("workbook", help="path to the .xlsx question bank")

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError("pip install openpyxl first") from exc
        wb = openpyxl.load_workbook(options["workbook"], data_only=True)

        rows = list(wb["Sections"].iter_rows(values_only=True))
        sections = [{"no": int(r[0]), "title": _s(r[1]), "intro": _s(r[2])} for r in rows[1:] if r[0] is not None]

        rows = list(wb["Answer Options"].iter_rows(values_only=True))
        head = [_s(h) for h in rows[0]]
        wording_col = next((i for i, h in enumerate(head) if h.startswith("Policy wording")), None)
        options_by_q: dict[str, list] = {}
        for r in rows[1:]:
            if not r[0]:
                continue
            wording = _s(r[wording_col]) if wording_col is not None else ""
            options_by_q.setdefault(_s(r[0]), []).append({
                "id": _s(r[2]),
                "text": _s(r[3]),
                "recommended": _s(r[4]).lower() == "yes",
                "note": _s(r[5]),
                # Full policy sentence for this answer (single-choice questions);
                # "(no sentence)" = the answer adds nothing.
                "wording": "" if wording == "(no sentence)" else wording,
                "no_sentence": wording == "(no sentence)",
            })

        rows = list(wb["Questions"].iter_rows(values_only=True))
        head = [_s(h) for h in rows[0]]
        questions = []
        for r in rows[1:]:
            if not r[2]:
                continue
            d = dict(zip(head, r))
            kind = TYPES.get(_s(d["Answer type"]))
            if kind is None:
                raise CommandError(f"Unknown answer type {d['Answer type']!r} for {d['Q ID']}")
            qid = _s(d["Q ID"])
            questions.append({
                "id": qid,
                "section": int(d["Section No"]),
                "text": _s(d["Question (shown to client)"]),
                "help": _s(d["Help text"]),
                "type": kind,
                "yes_no": _s(d["Answer type"]) == "Yes/No",
                "required": _s(d["Required"]).lower() == "yes",
                "show_if": _s(d["Show if"]),
                "allow_other": _s(d["Allow 'Other'"]).lower() == "yes",
                "fill_label": _s(d["Fill-in label"]),
                "example": _s(d["Example answer"]),
                "template": _s(d["Policy wording template"]),
                "options": options_by_q.get(qid, []),
            })

        OUT.write_text(json.dumps({"source": Path(options["workbook"]).name, "sections": sections,
                                   "questions": questions}, ensure_ascii=False, indent=1), encoding="utf-8")
        self.stdout.write(f"{len(sections)} sections, {len(questions)} questions, "
                          f"{sum(len(q['options']) for q in questions)} options -> {OUT}")
